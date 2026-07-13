# -*- coding: utf-8 -*-
"""スプレッドシート(新入社員一覧)の読み込み。

Googleスプレッドシートをxlsx形式でダウンロードしたもの、
または同じ列構成のExcelファイルを読み込む。
ヘッダーは2行目、データは3行目以降にある前提。
日付セルは「07/06」「7月6日」のような文字列・日付型・Excelシリアル値が
混在しているため、すべて datetime.date に正規化する。
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from io import BytesIO
from typing import Optional

import openpyxl

EXCEL_EPOCH = dt.date(1899, 12, 30)

HEADER_ROW = 2


@dataclass
class Person:
    furigana: str = ""
    name: str = ""
    eigyoubu: str = ""
    shozoku_code: str = ""
    shozoku_name: str = ""      # 所属名(店舗名)
    yakushoku: str = ""         # 役職
    grade: str = ""             # 等級
    nyusha_date: Optional[dt.date] = None
    email: str = ""
    tencho: str = ""
    bucho: str = ""
    birth: Optional[dt.date] = None
    zenshoku: str = ""
    tresen_start: Optional[dt.date] = None
    tresen_end: Optional[dt.date] = None
    kyukyu1: Optional[dt.date] = None
    kyukyu2: Optional[dt.date] = None
    haizoku_date: Optional[dt.date] = None
    first_shift: str = ""       # 初日勤務開始時間
    shataku: bool = False       # 社宅案内が必要(転居=〇)
    done: bool = False          # 行がグレーアウト=対応済み
    row: int = 0                # 元の行番号(表示用)
    warnings: list = field(default_factory=list)


def _to_date(value, default_year: int) -> Optional[dt.date]:
    """様々な形式の日付表現を date に変換する。変換できなければ None。"""
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)):
        # Excelシリアル値
        if 20000 < value < 60000:
            return EXCEL_EPOCH + dt.timedelta(days=int(value))
        return None
    s = str(value).strip()
    m = re.match(r"^(\d{4})[/年\-](\d{1,2})[/月\-](\d{1,2})日?$", s)
    if m:
        return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.match(r"^(\d{1,2})[/月](\d{1,2})日?$", s)
    if m:
        return dt.date(default_year, int(m.group(1)), int(m.group(2)))
    return None


def _is_grayed(cell) -> bool:
    """セルの背景がグレー系(=対応済みの印)かどうか。"""
    fill = cell.fill
    if fill is None or fill.patternType is None:
        return False
    if fill.patternType in ("lightGray", "gray125", "gray0625", "mediumGray", "darkGray"):
        return True
    if fill.patternType != "solid":
        return False
    rgb = getattr(fill.fgColor, "rgb", None)
    if not isinstance(rgb, str) or len(rgb) != 8:
        return False
    r, g, b = (int(rgb[i:i + 2], 16) for i in (2, 4, 6))
    # 彩度が低く(=無彩色に近い)、白でも黒でもない → グレー
    return max(r, g, b) - min(r, g, b) <= 24 and 100 <= max(r, g, b) <= 235


def _s(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_people(source, default_year: Optional[int] = None) -> list[Person]:
    """xlsxファイル(パスまたはbytes)から新入社員リストを読み込む。"""
    if default_year is None:
        default_year = dt.date.today().year
    if isinstance(source, (bytes, bytearray)):
        wb = openpyxl.load_workbook(BytesIO(source), data_only=True)
    else:
        wb = openpyxl.load_workbook(source, data_only=True)
    ws = wb.active

    headers = {}
    for idx, cell in enumerate(ws[HEADER_ROW]):
        if cell.value is None:
            continue
        key = str(cell.value).strip()
        if key == "公休" and key in headers:
            key = "公休2"
        headers[key] = idx

    def col(row, name):
        idx = headers.get(name)
        if idx is None or idx >= len(row):
            return None
        return row[idx].value

    people = []
    for row in ws.iter_rows(min_row=HEADER_ROW + 1):
        name = _s(col(row, "氏名"))
        if not name:
            continue
        p = Person(
            furigana=_s(col(row, "ふりがな")),
            name=name,
            eigyoubu=_s(col(row, "営業部")),
            shozoku_code=_s(col(row, "所属コード")),
            shozoku_name=_s(col(row, "所属名")),
            yakushoku=_s(col(row, "役職")),
            grade=_s(col(row, "等級")),
            nyusha_date=_to_date(col(row, "入社日"), default_year),
            email=_s(col(row, "メールアドレス")),
            tencho=_s(col(row, "店長")),
            bucho=_s(col(row, "営業部長")),
            birth=_to_date(col(row, "生年月日"), default_year),
            zenshoku=_s(col(row, "前職")),
            tresen_start=_to_date(col(row, "トレセン研修開始日"), default_year),
            tresen_end=_to_date(col(row, "トレセン研修終了日"), default_year),
            kyukyu1=_to_date(col(row, "公休"), default_year),
            kyukyu2=_to_date(col(row, "公休2"), default_year),
            haizoku_date=_to_date(col(row, "店舗配属日"), default_year),
            first_shift=_s(col(row, "初日勤務開始時間")),
            shataku=_s(col(row, "転居")) == "〇",
            done=_is_grayed(row[headers.get("氏名", 1)]),
            row=row[0].row,
        )
        if not p.nyusha_date:
            p.warnings.append("入社日が読み取れませんでした")
        if not p.email:
            p.warnings.append("メールアドレスが空です")
        if not p.grade:
            p.warnings.append("等級が空です")
        if not (p.tresen_end and p.kyukyu1 and p.kyukyu2 and p.haizoku_date):
            p.warnings.append("研修〜配属の日程が読み取れません(シートの日程列を確認してください)")
        people.append(p)
    return people
